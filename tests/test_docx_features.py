"""최신본 docx 시험 임포트(2026-10-06)에서 나온 보완: 표(#69), 제N부 들어가며(#70), 코멘트 작성자(#71), 걸린 문장 전문(#73). 작은 합성 docx로 확인한다."""

from __future__ import annotations

import re
from pathlib import Path

import pytest
from docx import Document
from mcp import Client
from docx.oxml.ns import qn

from magus.book.importer import import_manuscript
from magus.melchior.check import check_chapter
from magus.server import build_server
from magus.store import LocalStore
from test_web import server_proc  # noqa: F401 - 픽스처(tests/는 패키지가 아니다)

BIB = [
    "Clack, J. A. (2012). Gaining ground: The origin and evolution of tetrapods. Indiana University Press.",
    "Daeschler, E. B., Shubin, N. H., & Jenkins, F. A. (2006). A Devonian tetrapod-like fish. Nature, 440, 757–763.",
]


class Manuscript:
    """시험용 원고. pieces는 문자열 또는 (문자열, 코멘트 본문[, 작성자[, 날짜]])."""

    def __init__(self) -> None:
        self.doc = Document()

    def heading(self, text: str, level: int = 1) -> Manuscript:
        self.doc.add_paragraph(text, style=f"Heading {level}")
        return self

    def para(self, *pieces) -> Manuscript:
        p = self.doc.add_paragraph()
        for piece in pieces:
            if isinstance(piece, tuple):
                self.comment(p.add_run(piece[0]), *piece[1:])
            else:
                p.add_run(piece)
        return self

    def comment(self, runs, text: str, author: str = "편집자", date: str = "2026-10-02T09:00:00Z") -> None:
        c = self.doc.add_comment(runs, text=text, author=author, initials=author[:2])
        c._comment_elm.set(qn("w:date"), date)  # python-docx는 지금 시각을 넣는다. 판 사이 지문 시험을 위해 고정한다

    def table(self, rows: list[list[str]]):
        table = self.doc.add_table(rows=len(rows), cols=max(len(r) for r in rows))
        for r, row in enumerate(rows):
            for c, text in enumerate(row):
                table.cell(r, c).text = text
        return table

    def bibliography(self, lines: list[str] = BIB) -> Manuscript:
        self.heading("참고문헌")
        for line in lines:
            self.doc.add_paragraph(line)
        return self

    def save(self, path: Path) -> Path:
        path.parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(str(path))
        return path


def test_tables_are_markdown_with_citations_and_comments(store: LocalStore, tmp_path: Path):
    m = Manuscript().heading("제1장 물가")
    m.para("물고기는 바닥을 짚었다.")
    table = m.table([["지표", "값", "출처"], ["산소", "낮음", "Clack(2012)"], ["합친 칸", "", "a|b"]])
    table.cell(2, 0).merge(table.cell(2, 1))  # gridSpan=2
    m.comment(table.cell(1, 1).paragraphs[0].runs, "수치 확인 필요")
    m.para("다음 문단(Daeschler, Shubin & Jenkins, 2006).")
    import_manuscript(store, "alice", m.bibliography().save(tmp_path / "t.docx"), project_id="b")

    body = store.read_chapter("alice", "b", 1).content
    assert "| 지표 | 값 | 출처 |\n|---|---|---|\n| 산소 | 낮음 | Clack(2012) |\n| 합친 칸 |  | a\\|b |" in body
    assert "[표 1]" not in body
    # 표 안의 인용도 인용 게이트가 센다(최신본: 표 37개, 17,396자가 빠져 있었다)
    assert {g.key for g in check_chapter(store, "alice", "b", 1).groups} == {"clack2012", "daeschler2006"}
    [issue] = store.list_issues("alice", "b")
    assert issue.chapter_number == 1 and "— 원문: “낮음”" in issue.text


def test_empty_table_keeps_placeholder(store: LocalStore, tmp_path: Path):
    m = Manuscript().heading("제1장 물가")
    m.para("본문.")
    m.table([["", ""], ["", ""]])
    import_manuscript(store, "alice", m.bibliography().save(tmp_path / "e.docx"), project_id="b")
    assert "[표 1]" in store.read_chapter("alice", "b", 1).content


def _part_intro_manuscript(path: Path) -> Path:
    """실원고 모양: 부 제목은 서식 없는 줄, 장·부 들어가며·맺음은 제목 1."""
    m = Manuscript()
    m.para("제1부 항온").heading("제1장  따뜻한 피는 왜 비싼가").para("항온의 비용(Clack, 2012).")
    m.para("제2부 평지와 털").para("부의 제사(題詞).")
    m.heading("제2부 들어가며  인간은 왜 넓은 평지에 사는가")
    m.para("지리산 천왕봉에서 ", ("평지를 내려다본다", "수치 확인 필요"), "(Daeschler, Shubin & Jenkins, 2006).")
    m.heading("제2장  아르디는 왜 나무에서 내려왔나").para("나무 위의 삶.")
    m.heading("맺음  다음 맥박").para("맺는 글.")
    return m.bibliography().save(path)


def test_part_intro_is_a_chapter_of_its_part(store: LocalStore, tmp_path: Path):
    from magus.book.export import assemble_markdown
    from magus.book.outline import build_outline

    report = import_manuscript(store, "alice", _part_intro_manuscript(tmp_path / "p.docx"), project_id="b")
    # 진짜 부 제목이 남는다(최신본: "제2부 평지와 털"이 버려지고 부 이름이 "들어가며 …"가 됐었다)
    assert [(p.number, p.title) for p in store.list_parts("alice", "b")] == [(1, "항온"), (2, "평지와 털")]
    chapters = store.list_chapters("alice", "b")  # 읽는 순서: 부 들어가며는 그 부의 첫 장 앞
    assert [(c.number, c.display, c.title, c.part_number) for c in chapters] == [
        (1, "제1장", "따뜻한 피는 왜 비싼가", 1),
        (102, "제2부 들어가며", "인간은 왜 넓은 평지에 사는가", 2),
        (2, "제2장", "아르디는 왜 나무에서 내려왔나", 2),
        (90, "맺음", "다음 맥박", None),
    ]
    assert report.gaps == [] and not report.comments_outside
    [issue] = store.list_issues("alice", "b")
    assert issue.chapter_number == 102  # 장 밖 코멘트가 아니다
    assert {g.key for g in check_chapter(store, "alice", "b", 102).groups} == {"daeschler2006"}
    unsplit = (store.root / "users/alice/projects/b/original/unsplit.md").read_text(encoding="utf-8")
    assert "부의 제사" in unsplit and "지리산" not in unsplit  # 부 제사만 부 서문으로 남는다

    outline = build_outline(store, "alice", "b")
    assert "2장 + 부 들어가며 1·맺음 · 2부 · " in outline
    assert "### 제2부 평지와 털\n- 제2부 들어가며(102) 인간은 왜 넓은 평지에 사는가 · draft" in outline
    heads = [line for line in assemble_markdown(store, "alice", store.get_project("alice", "b")).splitlines() if line.startswith("#")]
    assert heads == [
        "# 제1부 항온", "## 제1장 따뜻한 피는 왜 비싼가",
        "# 제2부 평지와 털", "## 제2부 들어가며 인간은 왜 넓은 평지에 사는가", "## 제2장 아르디는 왜 나무에서 내려왔나",
        "## 맺음 다음 맥박",
    ]


@pytest.mark.anyio
async def test_issue_lines_name_part_intro(settings, store: LocalStore, tmp_path: Path):
    import_manuscript(store, settings.user_id, _part_intro_manuscript(tmp_path / "p.docx"), project_id="b")
    async with Client(build_server(settings, store)) as client:
        text = (await client.call_tool("book_issues", {"project_id": "b"})).content[0].text
        added = (await client.call_tool("book_issue_add", {"text": "도입 다듬기", "chapter_number": 102, "project_id": "b"})).content[0].text
        read = (await client.call_tool("book_read_chapter", {"number": 102, "project_id": "b"})).content[0].text
    assert "- I1 · open · 제2부 들어가며 ¶1 · 원고 코멘트(편집자) · fact" in text
    assert added.startswith("I2 열림 (b · 제2부 들어가며 · user)")
    assert '"label": "제2부 들어가며"' in read


# ---- #71 코멘트 작성자: 누가 말했는지 보존한다 -------------------------------------------------------


def test_author_kind_by_name():
    from magus.book.comments import author_kind

    assert [author_kind(n) for n in ("Claude", "ChatGPT 4o", "Microsoft Copilot", "편집자", "Kai Kim", "", None)] == [
        "ai", "ai", "ai", "person", "person", None, None,
    ]
    assert author_kind("우리 편집 도우미", ["편집 도우미"]) == "person"  # 이름 전체가 같아야 더한 목록이 맞는다
    assert author_kind("편집 도우미", ["편집 도우미"]) == "ai"


def _authored_manuscript(path: Path) -> Path:
    m = Manuscript().heading("제1장 물가")
    m.para(("데본기 물가(Daeschler, Shubin & Jenkins, 2006)", "인용 확인 필요", "Claude", "2026-10-02T01:02:03Z"), ".")
    m.para(("이 문단은 길다", "문단 줄이기", "편집자"), ".")
    return m.bibliography().save(path)


def test_comment_author_is_kept_in_file_and_index(store: LocalStore, tmp_path: Path):
    from magus.book.outline import build_outline

    report = import_manuscript(store, "alice", _authored_manuscript(tmp_path / "a.docx"), project_id="b")
    assert "작성자 Claude(AI) 1, 편집자 1" in report.summary()
    ai, person = store.list_issues("alice", "b")
    assert (ai.author, ai.author_kind, ai.comment_date, ai.by_ai) == ("Claude", "ai", "2026-10-02T01:02:03Z", True)
    assert (person.author, person.author_kind, person.by_ai) == ("편집자", "person", False)
    issues_md = (store.root / "users/alice/projects/b/issues.md").read_text(encoding="utf-8")
    assert "- author: Claude\n- author_kind: ai\n- comment_date: 2026-10-02T01:02:03Z" in issues_md
    store.reindex()
    assert [i.author_kind for i in store.list_issues("alice", "b")] == ["ai", "person"]
    assert "코멘트 유래 확인 요청 2건(citation 1 · general 1) · 그중 AI 확인 요청 1건(AI가 원고에 단 코멘트)" in build_outline(store, "alice", "b")


@pytest.mark.anyio
async def test_ai_comments_are_marked_in_tools(settings, store: LocalStore, tmp_path: Path):
    import_manuscript(store, settings.user_id, _authored_manuscript(tmp_path / "a.docx"), project_id="b")
    async with Client(build_server(settings, store)) as client:
        lines = (await client.call_tool("book_issues", {"project_id": "b"})).content[0].text
        read = (await client.call_tool("book_read_chapter", {"number": 1, "project_id": "b"})).content[0].text
    assert "- I1 · open · 1장 ¶1 · AI 확인 요청(Claude) · citation→daeschler2006" in lines
    assert "- I2 · open · 1장 ¶2 · 원고 코멘트(편집자) · general" in lines
    assert '"I1 ¶1 [AI 확인 요청] 인용 확인 필요' in read and '"I2 ¶2 문단 줄이기' in read


def test_web_marks_ai_comments(server_proc, tmp_path: Path):  # noqa: F811 - test_web의 픽스처
    import httpx2

    base, _, _ = server_proc({})
    docx_bytes = _authored_manuscript(tmp_path / "a.docx").read_bytes()
    with httpx2.Client(base_url=base, follow_redirects=True, timeout=120) as c:
        csrf = re.search(r'name="csrf" value="([0-9a-f]+)"', c.get("/login?next=/shelf").text).group(1)
        made = c.post("/shelf/manuscript", data={"csrf": csrf, "floor": "new"}, files={"file": ("원고.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
        page = made.text  # 올린 원고(#94 → #96)
        assert "1장 + " not in page and "올린 원고 · 1장 ·" in page  # outline과 같은 장 수 표기
        assert "열린 확인 요청 2건" in page
        issues = c.get(str(made.url).split("?")[0] + "?tab=issues").text
    assert "그중 🤖 AI 확인 요청 1건(AI가 원고에 단 코멘트입니다)" in issues
    assert '<td><span class="tag">🤖 AI 확인 요청(Claude)</span></td><td>문헌 확인</td>' in issues  # 문구는 그대로, 알약 태그
    assert "<td>원고 코멘트(편집자)</td><td>기타</td>" in issues


# ---- #73 걸린 문장 전문과 문단 번호 -------------------------------------------------------------------


def test_anchor_full_text_and_paragraph(store: LocalStore, tmp_path: Path):
    long = "가" * 150 + "(Clack, 2012)" + "나" * 60
    m = Manuscript().heading("제1장 물가").para("첫 문단.").heading("1.1 절", 2)
    m.para("셋째 문단의 ", (long, "인용 확인 필요", "Claude"), " 끝.")
    import_manuscript(store, "alice", m.bibliography().save(tmp_path / "a.docx"), project_id="b")
    [issue] = store.list_issues("alice", "b")
    assert issue.anchor == long  # 본문 표시는 120자로 잘리지만, 전문은 따로 둔다(최신본: 129건 중 82건이 잘렸다)
    assert "— 원문: “" + long[:120] + "”" in issue.text
    [group] = check_chapter(store, "alice", "b", 1).groups
    assert issue.paragraph == 3 and group.as_dict()["paragraphs"] == [3]  # 제목(##)도 문단 하나: ref_check와 같은 번호
    store.reindex()
    assert store.list_issues("alice", "b")[0].anchor == long
    assert f"- paragraph: 3\n- anchor: {long}" in (store.root / "users/alice/projects/b/issues.md").read_text(encoding="utf-8")


# ---- #77 제목과 부제를 따로 ---------------------------------------------------------------------------


def _titled_manuscript(path: Path, core_title: str | None, cover: list[tuple[str, str | None]]) -> Path:
    m = Manuscript()
    m.doc.core_properties.title = core_title or ""
    for text, style in cover:
        m.doc.add_paragraph(text, style=style)
    m.doc.add_paragraph("목차")
    m.heading("제1장 물가").para("본문.")
    return m.bibliography().save(path)


def test_title_subtitle_and_english_titles(store: LocalStore, tmp_path: Path):
    from magus.book.export import export_docx
    from magus.book.outline import build_outline

    path = _titled_manuscript(
        tmp_path / "t.docx",
        "인류가 선택한 진화, 시대가 선택한 도시 — 생존과 번영 그리고 그 이동에 관한 역사",
        [("명태는 수온에 묶이고", None), ("인류가 선택한 진화,\n시대가 선택한 도시", None), ("생존과 번영 그리고 그 이동에 관한 역사", None),
         ("Evolution Chosen by Humans, Cities Chosen by Eras", None),
         ("A History of Survival, Prosperity, and Where They Moved", None), ("물고기의 첫 숨에서", None), ("2026년 10월", None)],
    )
    report = import_manuscript(store, "alice", path, project_id="b")
    p = store.get_project("alice", "b")
    assert (p.title, p.subtitle, p.title_en, p.subtitle_en) == (
        "인류가 선택한 진화, 시대가 선택한 도시", "생존과 번영 그리고 그 이동에 관한 역사",
        "Evolution Chosen by Humans, Cities Chosen by Eras", "A History of Survival, Prosperity, and Where They Moved",
    )
    assert "- 부제: 생존과 번영 그리고 그 이동에 관한 역사 · 영문: Evolution Chosen by Humans" in report.summary()
    outline = build_outline(store, "alice", "b").splitlines()
    assert outline[0] == "프로젝트: 인류가 선택한 진화, 시대가 선택한 도시 (b)"
    assert outline[1].startswith("부제: 생존과 번영 그리고 그 이동에 관한 역사 · 영문: Evolution Chosen by Humans, Cities Chosen by Eras — A History")
    store.reindex()  # 파일(project.json)이 원본이다
    assert store.get_project("alice", "b").subtitle == "생존과 번영 그리고 그 이동에 관한 역사"
    exported = Document(export_docx(store, "alice", store.get_project("alice", "b")))
    assert [p.text for p in exported.paragraphs if p.style.name in ("Title", "Subtitle")] == [
        "인류가 선택한 진화, 시대가 선택한 도시", "생존과 번영 그리고 그 이동에 관한 역사",
    ]


def test_title_from_cover_styles_and_explicit_arguments(store: LocalStore, tmp_path: Path):
    path = _titled_manuscript(tmp_path / "s.docx", None, [("닿으려는 종", "Title"), ("물에서 뭍으로", "Subtitle")])
    import_manuscript(store, "alice", path, project_id="a")
    a = store.get_project("alice", "a")
    assert (a.title, a.subtitle, a.title_en) == ("닿으려는 종", "물에서 뭍으로", None)
    import_manuscript(store, "alice", path, project_id="c", title="다른 제목", subtitle="다른 부제")
    c = store.get_project("alice", "c")
    assert (c.title, c.subtitle) == ("다른 제목", "다른 부제")
    path = _titled_manuscript(tmp_path / "n.docx", None, [("그냥 문단", None)])
    import_manuscript(store, "alice", path, project_id="d")
    assert (store.get_project("alice", "d").title, store.get_project("alice", "d").subtitle) == ("n", None)  # 파일 이름
