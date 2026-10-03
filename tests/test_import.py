from __future__ import annotations

from pathlib import Path

import pytest

from magus.book.importer import import_manuscript
from magus.book.outline import TOKEN_BUDGET, build_outline, estimate_tokens
from magus.melchior.check import refresh_cite_counts
from magus.store import LocalStore
from magus.synthetic import build_manuscript_docx, build_manuscript_pdf


@pytest.fixture(scope="module")
def manuscript(tmp_path_factory) -> Path:
    return build_manuscript_docx(tmp_path_factory.mktemp("fx") / "닿으려는_종_합성.docx")


@pytest.fixture
def report(store: LocalStore, manuscript: Path):
    return import_manuscript(store, "alice", manuscript, project_id="synth", synthetic=True)


def test_chapters_parts_and_toc_dedup(store: LocalStore, report):
    chapters = store.list_chapters("alice", "synth")
    assert [(c.number, c.title, c.part_number) for c in chapters] == [
        (1, "물가의 망설임", 1),
        (2, "목이 생긴 날", 1),  # "제 2 장  목이 생긴 날" 공백 변형
        (3, "숨의 경제학", 1),  # "제3장: …" 콜론
        (4, "닿으려는 손", 2),  # "제4장" 다음 줄의 제목
        (5, "뭍의 첫 밤", 2),
        (6, "고래의 귀환", 3),
        (7, "닿으려는 종", 3),
    ]
    assert [(p.number, p.title) for p in store.list_parts("alice", "synth")] == [(1, "기원"), (2, "뭍"), (3, "다시 물로")]
    assert len(report.discarded) == 11  # 목차의 부 3 + 장 7 + 참고문헌 1
    assert report.gaps == [] and report.empty_chapters == []
    assert "합성 fixture" in report.summary()


def test_placeholders_and_subheadings(store: LocalStore, report):
    body = lambda n: store.read_chapter("alice", "synth", n).content  # noqa: E731
    assert "## 데본기의 물가" in body(1)
    assert "[표 1]" in body(2) and "[그림 1]" in body(3) and "[수식]" in body(4)
    assert (report.figures, report.tables, report.equations) == (1, 1, 1)
    assert "목차" not in body(1)


def test_comments_become_classified_issues(store: LocalStore, report):
    issues = store.list_issues("alice", "synth")
    assert len(issues) == 13 and all(i.origin == "import_comment" for i in issues)
    by_text = {i.text.split("\n")[0]: i for i in issues}
    assert by_text["확인 필요"].chapter_number == 1
    assert by_text["확인 필요"].ref_key == "daeschler2006"  # 코멘트 본문에 없으면 앵커에서
    assert "— 원문: “(Daeschler, Shubin & Jenkins, 2006)”" in by_text["확인 필요"].text
    assert by_text["인용 확인필요 — 쪽수 표기"].ref_key == "clack2012"
    assert by_text["Diamond 1997 확인 필요"].ref_key == "diamond1997"
    assert by_text["원문 대조: Gould 1989 p.45"].ref_key == "gould1989"
    assert by_text["이 두 문단 순서 바꾸기"].chapter_number == 5
    assert by_text["서문 톤 조정"].chapter_number is None  # 분할 못 한 구간
    assert by_text["서지 확인 필요: 번역서 연도"].chapter_number is None  # 참고문헌 구간
    assert dict(report.comments_outside) == {"unsplit": 1, "bib": 1}
    assert {i.kind for i in issues} == {"citation", "fact", "general"}
    assert by_text["서지 확인 필요: 번역서 연도"].kind == "fact"  # 앵커(참고문헌 줄)에 인용 표기가 없다


def test_bibliography_and_files(store: LocalStore, report, data_dir: Path):
    assert report.bib.lines == 30 and len(report.bib.created) == 30 and len(report.bib.failed) == 3
    assert report.bib.collisions and report.bib.collisions[0].startswith("hu2023 → hu2023-2")
    refs = {r.key: r for r in store.list_refs("alice")}
    assert all(r.status == "unverified" for r in refs.values())
    assert {"daeschler2006", "clack2012", "clack2002", "galor2022", "henrich2016", "romer1966", "김영민2020"} <= set(refs)
    assert set(refs["galor2022"].aliases) == {"갤러2022", "갤러2023"}
    assert store.resolve_ref("alice", "헨릭2016").key == "henrich2016"
    original = data_dir / "users/alice/projects/synth/original"
    assert {p.name for p in original.iterdir()} == {"닿으려는_종_합성.docx", "unsplit.md", "import-report.md"}
    assert "| I2 | 1 | citation | daeschler2006 | 확인 필요 |" in (original / "import-report.md").read_text(encoding="utf-8")
    assert (data_dir / "users/alice/shelf/melchior/daeschler2006/bib.json").exists()


def test_reimport_is_safe(store: LocalStore, report, manuscript: Path, data_dir: Path):
    with pytest.raises(ValueError, match="replace"):
        import_manuscript(store, "alice", manuscript, project_id="synth")
    again = import_manuscript(store, "alice", manuscript, project_id="synth2")
    assert len(again.bib.created) == 0 and len(again.bib.existing) == 30  # 책장은 사람 소유: 다시 만들지 않는다
    store.write_chapter("alice", "synth", 1, "사람이 고친 본문", "수정")
    replaced = import_manuscript(store, "alice", manuscript, project_id="synth", replace=True)
    assert replaced.project_id == "synth"
    assert store.read_chapter("alice", "synth", 1).content.startswith("데본기")
    trash = list((data_dir / "users/alice/trash").iterdir())
    assert len(trash) == 1 and (trash[0] / "versions/ch-01.v1.md").exists()  # 고친 것도 이전본도 휴지통에 남는다


def test_import_reindex_roundtrip(store: LocalStore, report):
    before = (store.list_chapters("alice", "synth"), store.list_issues("alice", "synth", None), store.list_refs("alice"))
    assert any(c.unverified_cite_count for c in before[0])
    store.reindex()
    refresh_cite_counts(store, "alice")  # 장별 미확인 인용 수는 파생값: 재생성 뒤 서버 시작·magus reindex가 다시 계산한다
    assert (store.list_chapters("alice", "synth"), store.list_issues("alice", "synth", None), store.list_refs("alice")) == before


def test_outline_of_imported_project(store: LocalStore, report):
    text = build_outline(store, "alice", "synth")
    assert estimate_tokens(text) <= TOKEN_BUDGET
    assert "- 코멘트 유래 미해결 13건(citation " in text
    assert "### 제2부 뭍" in text and "- 제4장 닿으려는 손 · draft" in text


def test_pdf_fallback_word_layout(store: LocalStore, tmp_path: Path):
    """실원고(Word '메모 포함' PDF)의 레이아웃을 흉내 낸 합성 PDF: 줄 단위 텍스트, 반복 머리말·쪽번호, 들여쓰기,
    단어 중간 줄바꿈, 오른쪽 말풍선, 내어쓰기 참고문헌, 그림, 서장."""
    pdf = build_manuscript_pdf(tmp_path / "원고.pdf")
    report = import_manuscript(store, "alice", pdf, project_id="frompdf")
    assert "코멘트는 PDF 텍스트" in report.summary()
    chapters = store.list_chapters("alice", "frompdf")
    assert [(c.number, c.display, c.title, c.part_number) for c in chapters] == [
        (0, "서장", "물가에서", None),  # 번호 없는 장
        (1, "제1장", "물가의 망설임", 1),
        (2, "제2장", "목이 생긴 날", 1),
    ]
    ch1 = store.read_chapter("alice", "frompdf", 1).content
    # 단어 중간 줄바꿈을 공백 없이 잇고, 머리말·쪽번호·말풍선은 본문에 없다
    assert "물고기가 있었다(Daeschler, Shubin & Jenkins, 2006)." in ch1
    assert "닿으려는 종" not in ch1 and "메모 포함" not in ch1 and "논지 요약" not in ch1
    assert "## 1.1 물고기는 왜 뭍으로 올라왔나" in ch1 and "[그림 1]" in ch1
    issues = store.list_issues("alice", "frompdf")
    assert [(i.chapter_number, i.kind, i.ref_key) for i in issues] == [
        (0, "citation", "마셜2015"),  # 여러 줄 말풍선을 이어 붙임
        (1, "citation", "daeschler2006"),
        (2, "general", None),
    ]
    assert issues[1].ref_id == "daeschler2006"
    refs = {r.key: r for r in store.list_refs("alice")}
    assert {"clack2012", "daeschler2006", "nunn2011", "田代和生1981"} <= set(refs)  # 내어쓰기 항목을 한 줄로 이음
    assert {"diamond1997", "galor2022", "keynes1930"} <= set(refs)  # 이어지는 한 줄짜리 항목을 하나씩 가름
    assert len(refs) == 7
    assert set(refs["田代和生1981"].aliases) == {"다시로1981", "다시로가즈이1981"}
    assert report.bib.failed == []


def test_rejects_unknown_format_and_headless(store: LocalStore, tmp_path: Path):
    (tmp_path / "x.txt").write_text("제1장", encoding="utf-8")
    with pytest.raises(ValueError, match="지원하지 않는"):
        import_manuscript(store, "alice", tmp_path / "x.txt")
    from docx import Document

    doc = Document()
    doc.add_paragraph("장 제목이 없는 원고")
    doc.save(str(tmp_path / "plain.docx"))
    with pytest.raises(ValueError, match="제N장"):
        import_manuscript(store, "alice", tmp_path / "plain.docx")
