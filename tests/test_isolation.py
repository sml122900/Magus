"""사용자 격리 점검(2026-10-06, 아버지 계정 추가 전). 기존 격리 시험이 다루지 않던 것을 더한다.

이미 있는 시험
- tests/test_search.py::test_every_tool_is_user_scoped — 모든 MCP 도구를 다른 사용자로 부르면 빈 결과(읽기·쓰기)
- tests/test_auth.py::test_remote_auth_rejects_and_isolates_users — 실제 HTTP와 토큰 두 개, 401·403, 원고·받은함·생각
- tests/test_web.py::test_remote_web_login_with_oauth — 다른 사용자의 받은함이 안 보이고 내려받기 링크는 403
- tests/test_store.py::test_other_user_sees_nothing, tests/test_inbox.py(받은함 파일 이름), test_web.py(서명 위조·만료)

여기서 더하는 것
- 두 사용자가 같은 프로젝트 id(book)·같은 문헌 key·같은 받은함 파일 이름을 써도 섞이지 않는다(FTS·벡터 검색, reindex 뒤 포함)
- 파일 이름 인자(../, 절대 경로, 다른 사용자 이름)로 다른 공간에 닿지 못한다
- 토큰 안의 클레임으로 사용자를 고르지 못한다(user_id는 sub → 허용 목록으로만)
- 허용 목록에서 지운 뒤(재시작)에는 기존 토큰은 403, 기존 웹 세션은 로그인으로 돌아간다. 세션 쿠키를 고친 것은 무효
- 웹의 /book·/issues·/checklist·/inbox·삭제·임포트·내보내기를 다른 사용자 세션으로 하지 못한다
"""

from __future__ import annotations

import hashlib
import json
import re
import shutil
import time
from pathlib import Path

import httpx2
import jwt
import pytest
from docx import Document
from mcp import Client
from starlette.requests import Request
from test_auth import ISSUER, _client, _config, _token, auth_server, jwks_url, keys  # noqa: F401 - 픽스처
from test_http import FIXTURES
from test_web import _csrf, _login, fake_issuer, server_proc  # noqa: F401 - 픽스처

from magus.auth import USER_ID_CLAIM, AuthConfig, ForbiddenGate, JwtVerifier
from magus.book.demo import seed_demo
from magus.config import Settings
from magus.search import hybrid_search
from magus.server import build_server
from magus.signing import sign
from magus.store import Idea, LocalStore, Passage, Ref, StyleEntry
from magus.web import SESSION_COOKIE, Web

TAGS = {"호흡": "짧음", "시점": "1인칭", "열기": "장면", "장치": [], "어미": "-다", "톤": "건조"}


def _hashes(root: Path) -> dict[str, str]:
    return {p.relative_to(root).as_posix(): hashlib.sha256(p.read_bytes()).hexdigest() for p in root.rglob("*") if p.is_file()}


def _fill(store: LocalStore, user: str, word: str, n: int) -> None:
    """같은 id·key·파일 이름으로, 내용만 사용자마다 다르게."""
    store.create_project(user, f"{word}의 책", project_id="book")
    store.write_chapter(user, "book", 1, f"{word}의 본문이다(Clack, 2012).", "x", title=f"{word}의 장")
    store.add_decision(user, "book", f"{word}의 결정")
    store.add_issue(user, "book", f"{word}의 미해결", chapter_number=1)
    store.put_ref(user, Ref(key="clack2012", status="verified", bib={"title": f"{word}의 서지"}),
                  passages=[Passage(f"{word}의 발췌 네발동물", "근거")])
    store.add_ideas(user, [Idea(id=f"2026100600000000{n}-aaaaaa", text=f"{word}의 생각 바다", kind="thought", author="me")])
    store.add_style(user, StyleEntry(id=f"2026100600000000{n}-bbbbbb", passage=f"{word}의 구절", reason=f"{word}의 짧은 호흡", tags=TAGS))
    store.put_inbox(user, "원고.docx", word.encode())


class _SameVector:
    """모든 글을 같은 벡터로: 벡터 검색이 user_id로만 갈리는지 보려고."""

    def embed(self, texts: list[str]) -> list[list[float]]:
        return [[1.0, 0.0, 0.0] for _ in texts]


def test_same_ids_in_two_users_never_mix(store: LocalStore, data_dir: Path):
    _fill(store, "alice", "앨리스", 1)
    _fill(store, "bob", "밥", 2)
    alice_files = _hashes(data_dir / "users" / "alice")
    for _ in range(2):  # 처음, 그리고 인덱스를 파일에서 다시 만든 뒤
        for user, word, other in (("alice", "앨리스", "밥"), ("bob", "밥", "앨리스")):
            seen = [
                store.get_project(user, "book").title,
                store.read_chapter(user, "book", 1).content,
                *(d.text for d in store.list_decisions(user, "book")),
                *(i.text for i in store.list_issues(user, "book")),
                str(store.get_ref(user, "clack2012").bib),
                *(p.text for p in store.get_passages(user, "clack2012")),
                *(i.text for i in store.list_ideas(user)),
                *(s.reason for s in store.list_styles(user)),
                store.inbox_path(user, "원고.docx").read_bytes().decode(),
            ]
            for query, lanes in (("네발동물", ["ref_picked"]), ("바다", ["idea_me"]), ("호흡", ["style"])):
                seen += [c.text for c in hybrid_search(store, user, query, lanes, 10, None)]  # FTS
                seen += [c.text for c in hybrid_search(store, user, query, lanes, 10, _SameVector())]  # FTS + 벡터
            assert all(word in text for text in seen), (user, seen)
            assert not any(other in text for text in seen), (user, seen)
        store.reindex()
    assert _hashes(data_dir / "users" / "alice") == alice_files


@pytest.mark.anyio
async def test_file_arguments_cannot_reach_other_spaces(store: LocalStore, data_dir: Path, settings: Settings):
    """원격과 같은 규칙(build_server 기본: 받은함 이름만). 앨리스가 밥의 받은함·책장·프로젝트에 닿는 길을 모두 막는다."""
    seed_demo(store, "alice")
    bob_inbox = data_dir / "users" / "bob" / "inbox"
    bob_inbox.mkdir(parents=True)
    shutil.copy(FIXTURES / "닿으려는_종_합성.docx", bob_inbox / "b.docx")
    (bob_inbox / "b.md").write_text("**사용자:** 밥의 생각\n\n**Claude:** 답\n", encoding="utf-8")
    bob_before = _hashes(data_dir / "users" / "bob")
    paths = ["../bob/inbox/b.docx", "..\\bob\\inbox\\b.docx", str((bob_inbox / "b.docx").resolve()), "bob/inbox/b.docx",
             "b.docx", "../../users/bob/inbox/b.docx", "C:b.docx"]
    async with Client(build_server(settings, store)) as c:
        for path in paths:
            assert (await c.call_tool("book_import", {"path": path, "project_id": "x"})).is_error, path
            assert (await c.call_tool("ref_add", {"bib": {"authors": "Clack, J.", "year": "2012", "title": "G"},
                                                  "passages": [{"text": "t", "reason": "r"}], "file_path": path})).is_error, path
            md = path.replace(".docx", ".md")
            assert (await c.call_tool("idea_add_conversation", {"path": md})).is_error, md
        for key in ("../../bob/shelf/melchior/x", "..", "bob/x"):  # 문헌 key는 책장 폴더 이름이 된다
            assert (await c.call_tool("ref_add", {"bib": {"key": key}, "passages": [{"text": "t", "reason": "r"}]})).is_error, key
        assert (await c.call_tool("book_import", {"path": "x.docx", "project_id": "../bob"})).is_error
        exported = (await c.call_tool("book_export", {"path": "../../../bob/projects/x.docx", "project_id": "demo"})).content[0].text
    assert Path(exported.splitlines()[0].removeprefix("docx를 만들었다: ")).parent == data_dir.resolve() / "users/alice/projects/demo/exports"
    assert store.list_projects("alice")[0].id == "demo" and len(store.list_projects("alice")) == 1
    assert _hashes(data_dir / "users" / "bob") == bob_before  # 밥의 공간은 한 바이트도 바뀌지 않았다


def test_token_claims_cannot_choose_the_user(keys):  # noqa: F811
    """user_id는 서명된 sub → 허용 목록으로만 정한다. 토큰에 넣은 user_id 같은 클레임은 쓰지 않는다."""
    good, _, jwks = keys
    verifier = JwtVerifier(_config(), jwks)
    now = int(time.time())
    claims = {"sub": "sub-alice", "aud": "http://magus.test/mcp", "iss": ISSUER, "iat": now, "exp": now + 600,
              USER_ID_CLAIM: "bob", "user_id": "bob", "uid": "bob"}
    token = jwt.encode(claims, good, algorithm="RS256", headers={"kid": "k1"})
    import asyncio

    access = asyncio.run(verifier.verify_token(token))
    assert access is not None and access.claims[USER_ID_CLAIM] == "alice"


def _gate_status(gate: ForbiddenGate, token: str) -> int:
    import asyncio

    sent: list[dict] = []

    async def app(scope, receive, send):
        await send({"type": "http.response.start", "status": 200, "headers": []})

    gate.app = app

    async def send(message):
        sent.append(message)

    scope = {"type": "http", "path": "/mcp", "headers": [(b"authorization", f"Bearer {token}".encode())]}
    asyncio.run(gate(scope, None, send))
    return sent[0]["status"]


def _session_user(web: Web, cookie: str) -> tuple[str, str] | None:
    return web.user_of(Request({"type": "http", "headers": [(b"cookie", f"{SESSION_COOKIE}={cookie}".encode())]}))


def test_removed_user_loses_token_and_session(keys, store: LocalStore, settings: Settings):  # noqa: F811
    """허용 목록은 시작할 때 읽는다. 지우고 다시 띄우면(up -d) 그 사용자의 유효한 토큰은 403, 웹 세션은 무효."""
    good, _, jwks = keys
    token = _token(good)  # sub-alice
    before = AuthConfig.from_env({"AUTH_ISSUER": ISSUER, "PUBLIC_URL": "http://magus.test", "ALLOWED_USERS": "sub-alice=alice,sub-bob=bob"})
    after = AuthConfig.from_env({"AUTH_ISSUER": ISSUER, "PUBLIC_URL": "http://magus.test", "ALLOWED_USERS": "sub-bob=bob"})
    assert _gate_status(ForbiddenGate(None, JwtVerifier(before, jwks)), token) == 200
    assert _gate_status(ForbiddenGate(None, JwtVerifier(after, jwks)), token) == 403
    secret = b"s" * 40
    cookie = sign({"uid": "alice", "sid": "abc"}, secret, "session", 3600)
    web_before = Web(store, settings, "http://magus.test", secret, _allowed={"alice", "bob"})
    web_after = Web(store, settings, "http://magus.test", secret, _allowed={"bob"})
    assert _session_user(web_before, cookie) == ("alice", "abc")
    assert _session_user(web_after, cookie) is None  # 서명은 맞지만 목록에 없다
    # 쿠키의 uid를 밥으로 고치면 서명이 맞지 않는다. 다른 용도(내려받기)로 서명한 토큰도 세션이 아니다
    body, mac = cookie.split(".")
    forged = json.loads(jwt.utils.base64url_decode(body))
    forged["uid"] = "bob"
    assert _session_user(web_before, jwt.utils.base64url_encode(json.dumps(forged).encode()).decode() + "." + mac) is None
    assert _session_user(web_before, sign({"uid": "bob", "sid": "x"}, secret, "download", 3600)) is None
    assert _session_user(web_before, sign({"uid": "bob", "sid": "x"}, b"t" * 40, "session", 3600)) is None


def _mini_docx(path: Path, comment: str) -> Path:
    doc = Document()
    doc.add_paragraph("제1장 밥의 장", style="Heading 1")
    p = doc.add_paragraph()
    doc.add_comment(p.add_run("밥의 문장이다."), text=comment, author="편집자", initials="ED")
    doc.add_paragraph("참고문헌", style="Heading 1")
    doc.add_paragraph("Bob, B. (2020). Bob's book. Press.")
    doc.save(str(path))
    return path


def test_web_pages_are_per_user_even_with_same_project_id(server_proc, fake_issuer, tmp_path: Path):  # noqa: F811
    issuer, issuer_state, _ = fake_issuer
    base, data_dir, _ = server_proc({
        "PUBLIC_URL": "http://127.0.0.1:{port}", "AUTH_ISSUER": issuer, "ALLOWED_USERS": "sub-alice=alice,sub-bob=bob",
        "SESSION_SECRET": "x" * 40, "WEB_CLIENT_ID": "web-client",
    })
    with httpx2.Client(base_url=base, timeout=120) as alice, httpx2.Client(base_url=base, timeout=120) as bob:
        assert _login(alice, issuer_state, "sub-alice").status_code == 303
        assert _login(bob, issuer_state, "sub-bob").status_code == 303
        a_csrf, b_csrf = _csrf(alice.get("/inbox").text), _csrf(bob.get("/inbox").text)
        alice.post("/inbox/upload", data={"csrf": a_csrf}, files={"files": ("원고.docx", (FIXTURES / "닿으려는_종_합성.docx").read_bytes())})
        alice.post("/book/import", data={"csrf": a_csrf, "name": "원고.docx", "project_id": "synth", "mode": "new"})
        # 밥: 앨리스의 받은함 파일을 임포트·삭제하지 못하고, 이름에 경로를 넣어 올리지도 못한다
        assert "받은함에" in bob.post("/book/import", data={"csrf": b_csrf, "name": "원고.docx", "project_id": "synth", "mode": "new"}).text
        for name in ("../alice/inbox/원고.docx", "..\\alice\\inbox\\원고.docx", "원고.docx"):
            bob.post("/inbox/delete", data={"csrf": b_csrf, "name": name})
        page = bob.post("/inbox/upload", data={"csrf": b_csrf}, files={"files": ("../alice/inbox/x.pdf", b"%PDF")}).text
        assert "파일 이름으로 쓸 수 없다" in page
        assert "x.pdf" not in [p.name for p in (data_dir / "users" / "alice" / "inbox").iterdir()]
        assert (data_dir / "users" / "alice" / "inbox" / "원고.docx").exists()
        assert "프로젝트가 없습니다" in bob.post("/book/export", data={"csrf": b_csrf, "project_id": "synth"}).text
        # 밥도 같은 id(synth)로 자기 원고를 올린다
        bob.post("/inbox/upload", data={"csrf": b_csrf}, files={"files": ("밥.docx", _mini_docx(tmp_path / "b.docx", "밥의 코멘트").read_bytes())})
        assert not re.search("class=\"msg err\"", bob.post("/book/import", data={"csrf": b_csrf, "name": "밥.docx", "project_id": "synth", "mode": "new"}).text)
        pages = {who: {path: c.get(f"{path}?project=synth").text for path in ("/book", "/issues", "/checklist", "/inbox")}
                 for who, c in (("alice", alice), ("bob", bob))}
    assert "닿으려는 종" in pages["alice"]["/book"] and "닿으려는 종" not in pages["bob"]["/book"]
    assert "인용 확인필요 — 쪽수 표기" in pages["alice"]["/issues"] and "밥의 코멘트" not in pages["alice"]["/issues"]
    assert "밥의 코멘트" in pages["bob"]["/issues"] and "인용 확인필요" not in pages["bob"]["/issues"]
    assert "Daeschler" in pages["alice"]["/checklist"] and "Daeschler" not in pages["bob"]["/checklist"]
    assert "원고.docx" in pages["alice"]["/inbox"] and "원고.docx" not in pages["bob"]["/inbox"] and "밥.docx" in pages["bob"]["/inbox"]
    assert "밥.docx" not in pages["alice"]["/inbox"]


def test_web_session_ends_when_user_is_removed(server_proc, fake_issuer):  # noqa: F811
    """같은 SESSION_SECRET으로 다시 띄웠는데 ALLOWED_USERS에서 앨리스를 뺐다: 앨리스의 세션 쿠키와 내려받기 링크가 통하지 않는다."""
    issuer, issuer_state, _ = fake_issuer
    env = {"PUBLIC_URL": "http://127.0.0.1:{port}", "AUTH_ISSUER": issuer, "SESSION_SECRET": "x" * 40, "WEB_CLIENT_ID": "web-client"}
    base1, _, _ = server_proc({**env, "ALLOWED_USERS": "sub-alice=alice,sub-bob=bob"})
    with httpx2.Client(base_url=base1, timeout=120) as alice, httpx2.Client(base_url=base1, timeout=60) as bob:
        _login(alice, issuer_state, "sub-alice")
        _login(bob, issuer_state, "sub-bob")
        cookie, bob_cookie = alice.cookies.get(SESSION_COOKIE), bob.cookies.get(SESSION_COOKIE)
        csrf = _csrf(alice.get("/inbox").text)
        alice.post("/inbox/upload", data={"csrf": csrf}, files={"files": ("원고.docx", (FIXTURES / "닿으려는_종_합성.docx").read_bytes())})
        alice.post("/book/import", data={"csrf": csrf, "name": "원고.docx", "project_id": "synth", "mode": "new"})
        link = alice.post("/book/export", data={"csrf": csrf, "project_id": "synth"}, follow_redirects=False).headers["location"]
    base2, _, _ = server_proc({**env, "ALLOWED_USERS": "sub-bob=bob"})
    with httpx2.Client(base_url=base2, timeout=60, cookies={SESSION_COOKIE: bob_cookie}) as kept:  # 대조: 목록에 남은 밥은 그대로
        assert kept.get("/inbox", follow_redirects=False).status_code == 200
    with httpx2.Client(base_url=base2, timeout=60, cookies={SESSION_COOKIE: cookie}) as removed:
        for path in ("/inbox", "/book", "/issues", "/checklist", link):
            response = removed.get(path, follow_redirects=False)
            assert response.status_code == 303 and response.headers["location"].startswith("/login"), path


@pytest.mark.anyio
async def test_concurrent_requests_keep_their_own_user(keys, auth_server):  # noqa: F811
    """두 사람이 동시에 써도 요청마다 사용자가 섞이지 않는다(토큰 → 요청 맥락 → 작업 스레드의 도구)."""
    import anyio

    good, _, _ = keys
    base, _ = auth_server
    url = f"{base}/mcp"
    tokens = {"앨리스": _token(good, aud=url), "밥": _token(good, aud=url, sub="sub-bob")}
    for word, token in tokens.items():
        async with _client(url, token) as c:
            await c.call_tool("idea_add", {"text": f"{word}의 바다 생각", "kind": "thought"})
    results: list[tuple[str, str]] = []

    async def ask(word: str) -> None:
        async with _client(url, tokens[word]) as c:
            for _ in range(5):
                results.append((word, (await c.call_tool("idea_search", {"query": "바다", "author": "all"})).content[0].text))

    async with anyio.create_task_group() as tg:
        for _ in range(4):
            tg.start_soon(ask, "앨리스")
            tg.start_soon(ask, "밥")
    assert len(results) == 40
    for word, text in results:
        other = "밥" if word == "앨리스" else "앨리스"
        assert f"{word}의 바다 생각" in text and f"{other}의 바다 생각" not in text, (word, text)


def test_shelf_pages_show_only_own_shelf(server_proc, fake_issuer):  # noqa: F811
    """영감·문헌·문체 화면도 세션 사용자의 책장만 보인다(같은 문헌 key·같은 project id여도)."""
    issuer, issuer_state, _ = fake_issuer
    base, data_dir, _ = server_proc({
        "PUBLIC_URL": "http://127.0.0.1:{port}", "AUTH_ISSUER": issuer, "ALLOWED_USERS": "sub-alice=alice,sub-bob=bob",
        "SESSION_SECRET": "x" * 40, "WEB_CLIENT_ID": "web-client",
    })
    store = LocalStore(data_dir)
    for n, (user, word) in enumerate((("alice", "앨리스"), ("bob", "밥"))):
        _fill(store, user, word, n + 1)
    store.close()
    pages = {}
    for user in ("alice", "bob"):
        with httpx2.Client(base_url=base, timeout=120) as c:
            assert _login(c, issuer_state, f"sub-{user}").status_code == 303
            pages[user] = {path: c.get(f"{path}?project=book").text for path in ("/ideas", "/refs", "/styles")}
    for user, word, other in (("alice", "앨리스", "밥"), ("bob", "밥", "앨리스")):
        for path, page in pages[user].items():
            assert word in page and f"{other}의" not in page, (user, path)
