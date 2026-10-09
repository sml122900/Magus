"""책 = 게시판 페이지(2026-10-09 사용자 요청; 첨부는 #97). 게시판 목록(제목·📎·미리보기·날짜, 최신 글 위, 20개씩, 글쓰기 하나),
글쓰기(제목·본문·첨부, 끌어다 놓기·고르기·사진 찍기), 글 보기(본문·그림·파일 목록), 위치 표시(프로젝트 › 노드 › 책),
게시판 첨부 문서가 files_list·book_import에 보이고 읽힌다, 다른 사용자의 책·글·첨부에는 닿지 못한다."""

from __future__ import annotations

import re
import urllib.parse

import httpx2
import pytest
from mcp import Client
from test_web import _csrf, _login, fake_issuer, server_proc  # noqa: F401 - 픽스처

from magus.inbox import FileResolver, attached_file
from magus.server import build_server
from magus.store import CategoryItem, LocalStore, Passage, Ref

PNG = b"\x89PNG\r\n\x1a\n" + b"0" * 64
NOTES = "notes-common"


def _seed(data_dir, user: str, word: str) -> None:
    store = LocalStore(data_dir)
    store.create_project(user, f"{word}의 책", project_id="b")
    store.write_chapter(user, "b", 1, f"{word}의 문장(Clack, 2012).", "x", title="하나")
    store.create_category(user, "메모", "blank", "#1f2f5c", (), "common", category_id=NOTES)
    store.create_category(user, "문헌", "blank", "#1f2f5c", (), "b", category_id="refs-b")
    store.put_ref(user, Ref(key="clack2012", status="unverified", bib={"title": "Gaining ground", "authors": "Clack, J."}, book="refs-b"),
                  passages=[Passage(f"{word}의 쓸 부분", "근거", "p.3")])
    store.put_category_item(user, CategoryItem("20261008000000000-aaaaaa", NOTES, f"{word}의 글", "me", reason="이유는 그대로"))
    from magus.melchior.check import refresh_cite_counts  # noqa: PLC0415

    refresh_cite_counts(store, user)
    store.close()


def test_board_write_view_and_attachments(server_proc):  # noqa: F811
    base, data_dir, _ = server_proc({"MAX_ATTACH_MB": "0.01"})  # 첨부 한 개 10KB까지(시험용)
    _seed(data_dir, "alice", "앨리스")
    with httpx2.Client(base_url=base, follow_redirects=True, timeout=120) as c:
        board = c.get(f"/login?next=/shelf/{NOTES}").text
        csrf = _csrf(board)
        # 게시판: 위치 표시(프로젝트 › 「내 서재」 › 노드 › 책), 오른쪽 위 글쓰기 하나, 옆 패널(iframe) 없음
        assert re.search(r'<nav class="trail" aria-label="위치"><a href="/shelf">프로젝트</a>.*?<a href="/shelf/p/common">.*?'
                         r'<span aria-current="page">메모</span></nav>', board)
        assert f'<a class="button write-btn" href="/shelf/write/{NOTES}">글쓰기</a>' in board
        assert "<iframe" not in board and 'enctype="multipart/form-data"' not in board
        write = c.get(f"/shelf/write/{NOTES}").text
        assert 'enctype="multipart/form-data"' in write and 'name="files" multiple accept=".pdf,' in write
        assert 'accept="image/*" capture="environment"' in write and "끌어다 놓기" in write and "data-dropzone" in write
        assert re.search(r'<a href="/shelf/notes-common">메모</a>.*?<span aria-current="page">글쓰기</span>', write)
        # 새 글: 제목·본문·첨부 둘 → 글 보기 페이지(그림은 글 안에, 다른 파일은 이름을 누르면 내려받기)
        made = c.post(f"/shelf/write/{NOTES}", data={"csrf": csrf, "title": "답사", "content": "시장 사진"},
                      files=[("files", ("C:\\fakepath\\시장.png", PNG, "image/png")), ("files", ("표.pdf", b"%PDF-1.4 x", "application/pdf"))])
        item = re.search(rf"/shelf/{NOTES}/(\d{{17}}-[0-9a-f]{{6}})$", str(made.url)).group(1)
        assert "<h1>답사</h1>" in made.text and "<p>시장 사진</p>" in made.text
        assert re.search(rf'<div class="pics one"><a href="/shelf/{NOTES}/{item}/files/[^"]+" [^>]*data-full>'
                         rf'<img src="/shelf/{NOTES}/{item}/thumbs/[^"]+" alt="시장.png" loading="lazy">', made.text)  # 썸네일, 누르면 원본
        assert '<div class="lightbox" role="dialog" aria-modal="true"' in made.text
        assert re.search(r'download>표.pdf</a> <small>PDF · 1 KB</small>', made.text)
        links = list(dict.fromkeys(re.findall(rf'href="(/shelf/{NOTES}/{item}/files/[^"]+)"', made.text)))
        assert len(links) == 2
        png = c.get(links[0], headers={"sec-fetch-dest": "document", "sec-fetch-mode": "navigate"})  # 그대로 연다(지도로 넘기지 않음)
        assert png.status_code == 200 and png.content == PNG and png.headers["content-type"] == "image/png"
        assert png.headers["x-content-type-options"] == "nosniff" and png.headers["content-disposition"].startswith("inline")
        # 첨부만 있는 글도 된다(제목 자리에 파일 이름)
        only = c.post(f"/shelf/write/{NOTES}", data={"csrf": csrf, "title": "", "content": ""},
                      files={"files": ("녹취.txt", "음성 받아쓰기".encode(), "text/plain")})
        assert "<h1>녹취</h1>" in only.text  # 제목이 비고 첨부만 있으면 첫 첨부 이름이 제목(#99-2)
        txt = c.get(re.search(rf'href="(/shelf/{NOTES}/[^"]+/files/[^"]+)"[^>]*>녹취.txt', only.text).group(1))
        assert txt.headers["content-disposition"].startswith("attachment") and "filename*=utf-8''" in txt.headers["content-disposition"].lower()
        # 받지 않는 형식·너무 큰 파일·많은 파일·빈 글은 아무것도 저장하지 않는다
        for files, code in (([("files", ("x.svg", b"<svg/>", "image/svg+xml"))], "type"),
                            ([("files", ("x.exe", b"MZ", "application/octet-stream"))], "type"),
                            ([("files", ("큰.pdf", b"0" * 20_000, "application/pdf"))], "big"),
                            ([("files", (f"{n}.txt", b"x", "text/plain")) for n in range(11)], "many")):
            bad = c.post(f"/shelf/write/{NOTES}", data={"csrf": csrf, "title": "막혀야 함", "content": "x"}, files=files)
            # 저장하지 않고 그 자리에서 다시(적은 제목·본문은 그대로, 2026-10-09 — 전에는 ?error=로 넘기며 비웠다)
            assert bad.status_code == 400 and '<p class="msg err">' in bad.text and 'value="막혀야 함"' in bad.text, code
        empty = c.post(f"/shelf/write/{NOTES}", data={"csrf": csrf, "title": " ", "content": ""})
        assert empty.status_code == 400 and "제목이나 본문을 적어 주세요." in empty.text
        assert c.post(f"/shelf/write/{NOTES}", data={"csrf": "틀림", "content": "x"}).status_code == 403
        # 게시판 목록: 최신 글이 위, 첨부 표시(그림이면 작은 미리보기 + 📎2), 제목이 없으면 앞부분을 흐리게
        board = c.get(f"/shelf/{NOTES}").text
        rows = re.findall(r'<li><a class="post" href="(/shelf/[^"]+)">(.*?)</a>', board)
        assert [href for href, _ in rows][:2] == [f"/shelf/{NOTES}/{re.search(r'[^/]+$', str(only.url)).group(0)}", f"/shelf/{NOTES}/{item}"]
        assert '<span class="pt">녹취</span>' in rows[0][1] and '📎' in rows[0][1]
        assert f'<img class="thumb" src="/shelf/{NOTES}/{item}/thumbs/' in rows[1][1] and '📎 2</span>' in rows[1][1]
        assert '<span class="pt">답사</span>' in rows[1][1] and '📎 1</span>' in rows[0][1]  # 하나여도 개수
        assert rows[-1][0] == f"/shelf/{NOTES}/20261008000000000-aaaaaa"  # 가장 오래된 글이 아래
        # 문헌 글도 같은 게시판에(제목은 서지), 글 보기에는 쓸 부분만(이유·체크는 보이지 않는다)
        refs = c.get("/shelf/refs-b").text
        assert 'href="/shelf/refs-b/clack2012"' in refs and "Gaining ground" in refs
        ref = c.get("/shelf/refs-b/clack2012").text
        assert "<blockquote>「앨리스의 쓸 부분」</blockquote>" in ref and "근거" not in ref and "원문 확인</label>" not in ref
        assert c.get(f"/shelf/{NOTES}/없는-글").status_code == 404
    store = LocalStore(data_dir)
    assert [a.name for a in store.list_attachments("alice", f"{NOTES}/{item}")] == ["시장.png", "표.pdf"]
    assert store.get_category_item("alice", NOTES, "20261008000000000-aaaaaa").reason == "이유는 그대로"
    assert store.get_category_item("alice", NOTES, item).author == "me"
    store.close()


def test_board_pages_twenty_per_page_newest_first(server_proc):  # noqa: F811
    base, data_dir, _ = server_proc({})
    _seed(data_dir, "alice", "앨리스")
    store = LocalStore(data_dir)
    for n in range(44):  # 시드 글 1 + 44 = 45개 → 20 · 20 · 5
        store.put_category_item("alice", CategoryItem(f"20261009{n:09d}-bbbbbb", NOTES, f"글 {n:02d}", "me", f"제목 {n:02d}", "author",
                                                      created_at=f"2099-01-01T00:{n:02d}:00.000Z"))  # 시드 글(지금 시각)보다 늘 나중
    store.close()
    with httpx2.Client(base_url=base, follow_redirects=True, timeout=120) as c:
        first = c.get(f"/login?next=/shelf/{NOTES}").text
        titles = re.findall(r'<span class="pt">(제목 \d+)</span>', first)
        assert titles == [f"제목 {n:02d}" for n in range(43, 23, -1)]  # 최신 20개
        assert re.search(r'<nav class="pages" aria-label="쪽"><span aria-current="page">1</span><a href="[^"]*page=2">2</a>'
                         r'<a href="[^"]*page=3">3</a><a href="[^"]*page=2">다음</a></nav>', first)
        last = c.get(f"/shelf/{NOTES}?page=3").text
        assert re.findall(r'<span class="pt">(제목 \d+)</span>', last) == [f"제목 {n:02d}" for n in range(3, -1, -1)]
        assert "앨리스의 글" in last and ">이전</a>" in last and ">다음</a>" not in last
        assert "제목 03" in c.get(f"/shelf/{NOTES}?page=99").text and "제목 43" in c.get(f"/shelf/{NOTES}?page=x").text  # 넘치면 끝 쪽, 틀리면 첫 쪽
        found = c.get(f"/shelf/{NOTES}?q=" + urllib.parse.quote("제목 07")).text
        assert "<mark>" in found and 'href="/shelf/notes-common/20261009000000007-bbbbbb"' in found


@pytest.mark.anyio
async def test_board_attachment_reaches_files_list_and_book_import(settings, data_dir, tmp_path):
    """수정본 흐름: 게시판 글에 docx·pdf를 첨부하면 files_list에 '첨부:<id>'로 보이고, book_import(path=그 이름)가 읽는다."""
    from docx import Document  # noqa: PLC0415

    _seed(data_dir, "alice", "앨리스")
    doc = Document()
    doc.add_heading("제1장 처음", level=1)
    doc.add_paragraph("고친 첫 장의 본문입니다.")
    path = tmp_path / "고친 원고.docx"
    doc.save(path)
    store = LocalStore(data_dir)
    post = store.put_category_item("alice", CategoryItem("20261009120000000-cccccc", NOTES, "수정본 올립니다", "me", "수정본", "author"))
    att = store.put_attachment("alice", f"{NOTES}/{post.id}", "고친 원고.docx", path.read_bytes())
    store.put_attachment("alice", f"{NOTES}/{post.id}", "사진.png", PNG)  # 도구가 읽는 형식이 아니면 보이지 않는다
    hit = attached_file(store, "alice", att.id)
    assert hit and hit[0].read_bytes() == path.read_bytes() and hit[1] == f"shelf/categories/{NOTES}/files/{post.id}/{att.id}/고친 원고.docx"
    assert attached_file(store, "bob", att.id) is None  # 다른 사용자는 같은 id로도 못 읽는다
    assert FileResolver(store)("alice", f"첨부:{att.id}")[0] == hit[0]
    async with Client(build_server(settings, store)) as c:  # 원격과 같이 이름만 받는다(local_files=False)
        listed = (await c.call_tool("files_list", {})).content[0].text
        imported = await c.call_tool("book_import", {"path": f"첨부:{att.id}", "project_id": "b", "new_edition": True, "force": True})
        missing = await c.call_tool("book_import", {"path": "첨부:20990101000000000-ffffff", "project_id": "b", "new_edition": True})
    assert f"첨부:{att.id} · 고친 원고.docx" in listed and "「메모」 책의 글 「수정본」" in listed and "사진.png" not in listed
    assert not imported.is_error and "E2" in imported.content[0].text and missing.is_error
    assert store.list_editions("alice", "b")[-1].file == "고친 원고.docx"
    assert store.read_chapter("alice", "b", 1).content.strip() == "고친 첫 장의 본문입니다."
    store.close()


def test_board_pages_are_per_user(server_proc, fake_issuer):  # noqa: F811
    """같은 책 id·글 id·문헌 key를 써도 남의 책·글·첨부는 보지도, 쓰지도 못한다. 남에게만 있는 책은 없는 책이다."""
    issuer, issuer_state, _ = fake_issuer
    base, data_dir, _ = server_proc({
        "PUBLIC_URL": "http://127.0.0.1:{port}", "AUTH_ISSUER": issuer, "ALLOWED_USERS": "sub-alice=alice,sub-bob=bob",
        "SESSION_SECRET": "x" * 40, "WEB_CLIENT_ID": "web-client",
    })
    _seed(data_dir, "alice", "앨리스")
    _seed(data_dir, "bob", "밥")
    store = LocalStore(data_dir)
    store.create_category("alice", "앨리스만", "blank", "#1f2f5c", (), "common", category_id="alice-only")
    hidden = store.put_category_item("alice", CategoryItem("20261009000000000-dddddd", "alice-only", "앨리스만의 글", "me"))
    secret = store.put_attachment("alice", f"{NOTES}/20261008000000000-aaaaaa", "비밀.pdf", b"%PDF alice-only")
    ref_file = store.put_attachment("alice", "ref:clack2012", "원문.pdf", b"%PDF alice-ref")
    store.close()
    before = {p: p.read_bytes() for p in (data_dir / "users" / "alice").rglob("*") if p.is_file()}
    url = f"/shelf/{NOTES}/20261008000000000-aaaaaa/files/{secret.id}"
    ref_url = f"/shelf/refs-b/clack2012/files/{ref_file.id}"
    with httpx2.Client(base_url=base, timeout=120) as bob:
        assert bob.get(url, follow_redirects=False).status_code in (303, 307)  # 로그인 전: 로그인으로
        assert bob.get(f"/shelf/{NOTES}", follow_redirects=False).status_code in (303, 307)
        assert _login(bob, issuer_state, "sub-bob").status_code == 303
        csrf = _csrf(bob.get("/shelf").text)
        assert bob.get(url).status_code == 404 and bob.get(ref_url).status_code == 404
        assert bob.get("/shelf/alice-only").status_code == 404 and bob.get(f"/shelf/alice-only/{hidden.id}").status_code == 404
        assert bob.get("/shelf/write/alice-only").status_code == 404
        assert bob.post("/shelf/write/alice-only", data={"csrf": csrf, "content": "밥이 씀"}, follow_redirects=False).status_code == 404
        own = bob.get(f"/shelf/{NOTES}/20261008000000000-aaaaaa").text  # 같은 id의 글은 밥의 것
        assert "밥의 글" in own and "앨리스" not in own and "비밀.pdf" not in own
        assert "비밀.pdf" not in bob.get(f"/shelf/{NOTES}").text and "원문.pdf" not in bob.get("/shelf/refs-b/clack2012").text
        for path in (url, ref_url, f"{url}/../../edit", "/shelf/notes-common/..%2F..%2Ffather/files/x", "/shelf/..%2F..%2Falice/notes-common"):
            got = bob.get(urllib.parse.quote(path, safe="/%"))  # ..는 밥의 공간 안으로만 풀린다
            assert b"alice" not in got.content and "앨리스".encode() not in got.content, path
    after = {p: p.read_bytes() for p in (data_dir / "users" / "alice").rglob("*") if p.is_file()}
    assert before == after
