"""3단계: book_decide / book_issue_add / book_issue_close / book_export."""

from __future__ import annotations

import dataclasses
import re
from pathlib import Path

import docx
import pytest
from mcp import Client

from magus.book.demo import seed_demo
from magus.book.importer import import_manuscript
from magus.config import Settings
from magus.server import build_server
from magus.store import LocalStore
from magus.synthetic import build_manuscript_docx

pytestmark = pytest.mark.anyio


def _text(result) -> str:
    assert not result.is_error, result.content[0].text
    return "".join(block.text for block in result.content)


async def test_decide_issue_add_close(settings: Settings, store: LocalStore, data_dir: Path):
    seed_demo(store, settings.user_id)
    async with Client(build_server(settings, store)) as client:
        assert _text(await client.call_tool("book_decide", {"text": "3장은 폐의 기원으로 시작한다."})).startswith("D3 기록함")
        added = _text(await client.call_tool("book_issue_add", {"text": "4장 연결 문단 필요", "chapter_number": 4, "origin": "agent"}))
        assert added.startswith("I4 열림 (demo · 4장 · agent)")
        assert (await client.call_tool("book_issue_add", {"text": "x", "chapter_number": 99})).is_error
        outline = _text(await client.call_tool("book_outline", {}))
        assert "- D3 " in outline and "- I4 · 4장 · agent · 4장 연결 문단 필요" in outline

        assert _text(await client.call_tool("book_issue_close", {"id": "I2", "note": "출처 찾음: Clack 2012"})) == "I2 닫음: 출처 찾음: Clack 2012"
        assert _text(await client.call_tool("book_issue_close", {"id": "4"})) == "I4 닫음"
        closed = _text(await client.call_tool("book_issues", {"status": "closed"}))
        assert "- I2 · closed · 3장 · user · general · 3장 폐의 기원 출처 찾기 · 닫음: 출처 찾음: Clack 2012" in closed
        assert (await client.call_tool("book_issue_close", {"id": "I99"})).is_error
        assert (await client.call_tool("book_issue_close", {"id": "삼"})).is_error
        # 닫은 번호는 다시 쓰지 않는다
        assert _text(await client.call_tool("book_issue_add", {"text": "새 이슈"})).startswith("I5 열림")
    issues_md = (data_dir / "users/alice/projects/demo/issues.md").read_text(encoding="utf-8")
    assert "## I2 · closed" in issues_md and "- closed_note: 출처 찾음: Clack 2012" in issues_md


async def test_instructions_rule4_active(settings: Settings, store: LocalStore):
    async with Client(build_server(settings, store)) as client:
        assert "book_write_chapter는 저자가 Magus 원고에 저장하라고 명시했을 때만 쓰고" in client.instructions
        tools = {t.name: t for t in (await client.list_tools()).tools}
        assert tools["book_write_chapter"].description.startswith("저자가 'Magus 원고에 저장'하라고 분명히 말했을 때만 부른다.")


def _md_blocks(markdown: str) -> list[str]:
    return [b.strip() for b in re.split(r"\n\s*\n", markdown) if b.strip()]


async def test_export_roundtrip_titles_and_paragraphs(settings: Settings, store: LocalStore, tmp_path: Path):
    import_manuscript(store, settings.user_id, build_manuscript_docx(tmp_path / "m.docx"), project_id="synth")
    async with Client(build_server(settings, store)) as client:
        result = _text(await client.call_tool("book_export", {"path": "원고.docx", "project_id": "synth"}))
    location = result.splitlines()[0].removeprefix("docx를 만들었다: ")
    assert location.endswith("원고.docx") and Path(location).exists()
    paragraphs = docx.Document(location).paragraphs
    headings = [(p.style.name, p.text) for p in paragraphs if p.style.name.startswith("Heading")]
    chapters = store.list_chapters(settings.user_id, "synth")
    chapter_heads = [t for s, t in headings if s == "Heading 2"]
    assert chapter_heads == [f"제{c.number}장 {c.title}" for c in chapters]
    assert [t for s, t in headings if s == "Heading 1"] == ["제1부 기원", "제2부 뭍", "제3부 다시 물로"]
    # 장마다 docx 단락 수 == 장 본문 md 블록 수(소제목 포함, 표 블록은 docx에서 단락이 아니라 표가 된다)
    starts = [i for i, p in enumerate(paragraphs) if p.style.name == "Heading 2"]
    for n, chapter in enumerate(chapters):
        end = next((i for i in range(starts[n] + 1, len(paragraphs)) if paragraphs[i].style.name in ("Heading 1", "Heading 2")), len(paragraphs))
        body = store.read_chapter(settings.user_id, "synth", chapter.number).content
        assert end - starts[n] - 1 == len([b for b in _md_blocks(body) if not b.startswith("|")]), chapter.number
    assert any(p.text == "데본기의 물가" and p.style.name == "Heading 3" for p in paragraphs)
    # 표(#69)는 markdown 표로 들어갔다가 Word 표로 나간다
    [table] = docx.Document(location).tables
    assert [[c.text for c in row.cells] for row in table.rows] == [["구분", "숨쉬기"], ["물고기", "아가미"], ["네발동물", "폐"]]


async def test_export_and_tools_respect_user(settings: Settings, store: LocalStore):
    seed_demo(store, settings.user_id)
    bob = dataclasses.replace(settings, user_id="bob")
    async with Client(build_server(bob, store)) as client:
        for name, args in [
            ("book_decide", {"text": "x", "project_id": "demo"}),
            ("book_issue_add", {"text": "x", "project_id": "demo"}),
            ("book_issue_close", {"id": "I1", "project_id": "demo"}),
            ("book_export", {"project_id": "demo"}),
        ]:
            assert (await client.call_tool(name, args)).is_error, name
