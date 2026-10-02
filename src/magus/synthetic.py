"""합성 fixture. 실제 원고(fixtures/)가 없을 때 2·4·5단계를 돌리기 위한 가짜 원고·대화 기록을 만든다.

내용은 모두 지어낸 것이다(문헌 목록의 일부 서지는 실존 문헌을 흉내 냈지만 확인된 서지가 아니다).
`magus synth <폴더>`로 파일을 만들고 `magus import <폴더>/닿으려는_종_합성.docx`로 임포트한다.
"""

from __future__ import annotations

import io
from pathlib import Path

import pymupdf
from docx import Document
from docx.oxml import parse_xml
from docx.text.paragraph import Paragraph

MANUSCRIPT_NAME = "닿으려는_종_합성.docx"

TOC = [
    "제1부 기원",
    "제1장 물가의 망설임\t3",
    "제 2 장  목이 생긴 날\t9",
    "제3장: 숨의 경제학\t15",
    "제2부  뭍",
    "제4장 닿으려는 손\t21",
    "제5장 뭍의 첫 밤\t27",
    "제3부 다시 물로",
    "제6장 고래의 귀환\t33",
    "제7장 닿으려는 종\t39",
    "참고문헌\t45",
]

BIBLIOGRAPHY = [
    "Ahlberg, P. E. (2019). Follow the footprints and mind the gaps: A new look at the origin of tetrapods. Earth and Environmental Science Transactions, 109(1–2), 115–137.",
    "Benton, M. J. (2015). Vertebrate palaeontology (4th ed.). Wiley-Blackwell.",
    "Carroll, R. L. (1988). Vertebrate paleontology and evolution. W. H. Freeman.",
    "Clack, J. A. (2002). An early tetrapod from Romer's Gap. Nature, 418(6893), 72–76.",
    "Clack, J. A. (2012). Gaining ground: The origin and evolution of tetrapods (2nd ed.). Indiana University Press.",
    "Coates, M. I., & Clack, J. A. (1990). Polydactyly in the earliest known tetrapod limbs. Nature, 347(6288), 66–69.",
    "Daeschler, E. B., Shubin, N. H., & Jenkins, F. A. (2006). A Devonian tetrapod-like fish and the evolution of the tetrapod body plan. Nature, 440(7085), 757–763.",
    "Darwin, C. (1859). On the origin of species by means of natural selection. John Murray. [국역: 찰스 다윈, 『종의 기원』, 장대익 옮김, 사이언스북스, 2019]",
    "Dawkins, R. (1976). The selfish gene. Oxford University Press. [국역: 리처드 도킨스, 『이기적 유전자』, 홍영남·이상임 옮김, 을유문화사, 2018]",
    "Diamond, J. (1997). Guns, germs, and steel: The fates of human societies. W. W. Norton. [국역: 재레드 다이아몬드, 『총, 균, 쇠』, 김진준 옮김, 문학사상사, 1998]",
    "Galor, O. (2022). The journey of humanity: The origins of wealth and inequality. Dutton. [국역: 오데드 갤러, 『인류의 여정』, 강영옥 옮김, 시공사, 2023]",
    "Gould, S. J. (1989). Wonderful life: The Burgess Shale and the nature of history. W. W. Norton.",
    "Harari, Y. N. (2015). Sapiens: A brief history of humankind. Harper. [국역: 유발 하라리, 『사피엔스』, 조현욱 옮김, 김영사, 2015]",
    "Henrich, J. (2016). The secret of our success. Princeton University Press. [국역: 조지프 헨릭, 『호모 사피엔스, 그 성공의 비밀』, 주명진·이병권 옮김, 뿌리와이파리, 2019]",
    "Hu, Y., Lin, Q., & Zhang, R. (2023). A synthetic study of lobe-fin locomotion. Journal of Synthetic Paleontology, 12(3), 45–67.",
    "Hu, Z. (2023). Another synthetic study with the same surname and year. Synthetic Letters, 4, 1–9.",
    "Long, J. A. (2011). The rise of fishes: 500 million years of evolution (2nd ed.). Johns Hopkins University Press.",
    "Niedźwiedzki, G., Szrek, P., Narkiewicz, K., Narkiewicz, M., & Ahlberg, P. E. (2010). Tetrapod trackways from the early Middle Devonian period of Poland. Nature, 463(7277), 43–48.",
    "Pierce, S. E., Clack, J. A., & Hutchinson, J. R. (2012). Three-dimensional limb joint mobility in the early tetrapod Ichthyostega. Nature, 486(7404), 523–526.",
    "Romer, A. S. Vertebrate paleontology. University of Chicago Press, 1966.",
    "Shubin, N. (2008). Your inner fish. Pantheon. [국역: 닐 슈빈, 『내 안의 물고기』, 김명남 옮김, 김영사, 2009]",
    "Wilson, E. O. (1998). Consilience: The unity of knowledge. Knopf. [국역: 에드워드 윌슨, 『통섭』, 최재천·장대익 옮김, 사이언스북스, 2005]",
    "World Health Organization. (2020). A synthetic report on nothing in particular. WHO.",
    "Zimmer, C. (1998). At the water's edge. Free Press.",
    "김영민 (2020). 『아침에는 죽음을 생각하는 것이 좋다』. 사회평론.",
    "이정모 (2018). 『저도 과학은 어렵습니다만』. 바틀비.",
    "최재천·김산하 (2014). 『통섭의 식탁』. 움직이는서재.",
    "편집부. 『자연사 연표』. 출판사 미상.",
    "Smith, J. Unpublished manuscript on fins.",
    "——— 앞의 책, 23쪽.",
]

# ref_check용: 스펙의 인용 패턴과 4단계 보강 대상(괄호 안 다중 인용, 조사 붙은 한글 저자)
CITATION_PARAGRAPHS = [
    "데본기 얕은 물가에서 지느러미로 바닥을 짚는 물고기가 있었다(Daeschler, Shubin & Jenkins, 2006).",
    "Clack(2012)은 이 전이를 길게 다룬다. Hu et al.(2023)도 비슷한 결론에 이른다.",
    "갤러는(2022) 이것을 인류사 전체로 넓힌다. 협력은 문화로 축적되었다(헨릭, 2016).",
    "여러 연구가 같은 방향을 가리킨다(Clack, 2012; Hu et al., 2023). 갤러(2022)의 말처럼 문은 늘 좁았다.",
]


def _png() -> bytes:
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 8, 8), False)
    pix.clear_with(200)
    return pix.tobytes("png")


def _para(doc, *pieces, style: str | None = None) -> Paragraph:
    """pieces: 문자열 또는 (문자열, 코멘트). 코멘트가 있는 조각에 코멘트 범위를 건다."""
    p = doc.add_paragraph(style=style)
    for piece in pieces:
        if isinstance(piece, tuple):
            text, comment = piece
            run = p.add_run(text)
            doc.add_comment(run, text=comment, author="편집자", initials="ED")
        else:
            p.add_run(piece)
    return p


def build_manuscript_docx(path: Path) -> Path:
    doc = Document()
    doc.core_properties.title = "닿으려는 종 (합성)"
    doc.add_paragraph("닿으려는 종", style="Title")
    doc.add_paragraph("목차")
    for line in TOC:
        doc.add_paragraph(line)
    doc.add_paragraph("들어가며", style="Heading 1")
    _para(doc, ("이 책은 물에서 뭍으로, 다시 물로 돌아간 생물들의 이야기다.", "서문 톤 조정"), " 모든 장은 하나의 장면으로 시작한다.")

    doc.add_paragraph("제1부 기원", style="Heading 1")
    doc.add_paragraph("제1장 물가의 망설임", style="Heading 2")
    _para(doc, "데본기 얕은 물가에서 지느러미로 바닥을 짚는 물고기가 있었다", ("(Daeschler, Shubin & Jenkins, 2006)", "확인 필요"), ".")
    _para(doc, ("Clack(2012)", "인용 확인필요 — 쪽수 표기"), "은 이 전이를 길게 다룬다. Hu et al.(2023)도 비슷한 결론에 이른다.")
    doc.add_paragraph("데본기의 물가", style="Heading 3")
    _para(doc, ("갤러는(2022)", "원문 대조 필요"), " 이것을 인류사 전체로 넓힌다. 협력은 문화로 축적되었다(헨릭, 2016).")
    _para(doc, CITATION_PARAGRAPHS[3])
    _para(doc, ("이 문단은 지나치게 길다. 물가의 장면을 더 길게 묘사하고 싶지만 독자가 지칠 수 있다.", "이 문단은 너무 길다"))

    doc.add_paragraph("제 2 장  목이 생긴 날", style="Heading 2")
    _para(doc, "아가미덮개가 사라지고 목이 생겼다. ", ("물고기가 처음으로 고개를 돌렸다.", "물고기가 일어서지 못하는 장면 보강"))
    table = doc.add_table(rows=3, cols=2)
    for r, (a, b) in enumerate([("구분", "숨쉬기"), ("물고기", "아가미"), ("네발동물", "폐")]):
        table.cell(r, 0).text, table.cell(r, 1).text = a, b
    _para(doc, "목은 머리를 돌리는 자유였다. ", ("Hu et al. 2023 참고", "서지 정보 누락: Hu et al. 2023"))

    doc.add_paragraph("제3장: 숨의 경제학", style="Heading 2")
    _para(doc, "폐는 처음부터 육지를 위해 생긴 기관이 아니었다. [출처 필요: 초기 폐의 기능]")
    doc.add_picture(io.BytesIO(_png()))
    _para(doc, ("그림: 물가의 산소 농도", "그림 캡션 다듬기"))

    doc.add_paragraph("제2부  뭍", style="Heading 1")
    doc.add_paragraph("제4장", style="Heading 2")
    doc.add_paragraph("닿으려는 손")
    _para(doc, "손가락은 걷기보다 먼저 물풀을 헤치기 위해 생겼을지도 모른다. ", ("다이아몬드(1997)", "Diamond 1997 확인 필요"), "는 다른 맥락에서 같은 질문을 던진다.")
    eq = doc.add_paragraph("숨의 비용은 대략 다음과 같다: ")
    eq._p.append(parse_xml('<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"><m:r><m:t>C=kV</m:t></m:r></m:oMath>'))

    doc.add_paragraph("제5장 뭍의 첫 밤", style="Heading 2")
    first = doc.add_paragraph().add_run("첫 문단: 뭍의 밤은 차가웠다.")
    doc.add_paragraph("가운데 문단: 별이 물에 비치지 않았다.")
    last = doc.add_paragraph().add_run("마지막 문단: 그래도 다음 날 아침이 왔다.")
    doc.add_comment([first, last], text="이 두 문단 순서 바꾸기", author="저자", initials="AU")

    doc.add_paragraph("제3부 다시 물로", style="Heading 1")
    doc.add_paragraph("제6장 고래의 귀환", style="Heading 2")
    _para(doc, "고래는 뭍에서 다시 물로 돌아갔다. ", ("우연은 역사의 다른 이름이다.", "원문 대조: Gould 1989 p.45"))
    doc.add_paragraph("제7장 닿으려는 종", style="Heading 2")
    _para(doc, ("닿으려는 것은 종의 본성이다.", "결론 문단 다시 쓰기"), " 우리는 아직 닿지 않았다.")

    doc.add_paragraph("참고문헌", style="Heading 1")
    for line in BIBLIOGRAPHY:
        if line.startswith("Galor"):
            _para(doc, (line, "서지 확인 필요: 번역서 연도"))
        else:
            doc.add_paragraph(line)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path


def build_manuscript_pdf(path: Path) -> Path:
    """PDF 폴백 확인용. Word가 말풍선을 '메모 포함[C1]: …'로 내보낸 모양을 흉내 낸다."""
    lines = [
        "닿으려는 종",
        "제1부 기원",
        "제1장 물가의 망설임",
        "데본기 얕은 물가에서 지느러미로 바닥을 짚는 물고기가 있었다(Daeschler, Shubin & Jenkins, 2006).",
        "메모 포함[C1]: 확인 필요",
        "Clack(2012)은 이 전이를 길게 다룬다.",
        "제2장 목이 생긴 날",
        "아가미덮개가 사라지고 목이 생겼다.",
        "메모 포함[C2]: 이 문단은 너무 길다",
        "참고문헌",
        "Clack, J. A. (2012). Gaining ground. Indiana University Press.",
    ]
    doc = pymupdf.open()
    page = doc.new_page()
    y = 60
    for line in lines:
        page.insert_text((50, y), line, fontname="korea", fontsize=10)
        y += 30  # 줄마다 별도 블록이 되도록 간격을 둔다
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path


def write_all(folder: Path) -> dict[str, Path]:
    folder.mkdir(parents=True, exist_ok=True)
    return {
        "manuscript_docx": build_manuscript_docx(folder / MANUSCRIPT_NAME),
        "manuscript_pdf": build_manuscript_pdf(folder / "닿으려는_종_합성.pdf"),
    }
